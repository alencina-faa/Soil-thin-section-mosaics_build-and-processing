import math
import os

import h5py
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd
import seaborn as sns
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from scipy import stats
from sklearn.ensemble import RandomForestClassifier
from sklearn.metrics import roc_auc_score


# ==========================================
# 1. FUNCIÓN DE LECTURA Y EXTRACCIÓN .H5
# ==========================================
def process_h5_file(file_path):
    """
    Lee un archivo .h5, extrae Shape y Convex Shape de los poros interiores (is_edge==False),
    y calcula Irregularity = atan(Shape / Convex Shape).
    """
    data = []

    try:
        with h5py.File(file_path, "r") as f:
            if "contours" not in f:
                return None

            contours_group = f["contours"]

            for contour_id, group in contours_group.items():
                # 1. Filtrar únicamente contornos interiores (is_edge = False)
                is_edge_value = np.asarray(group.attrs.get("is_edge", False))
                is_edge = bool(is_edge_value.reshape(-1)[0]) if is_edge_value.size else False
                if is_edge:
                    continue

                record = {"Contour_ID": contour_id}
                for key, value in group.attrs.items():
                    if key == "is_edge":
                        continue
                    scalar = np.asarray(value)
                    if scalar.size == 1 and np.issubdtype(scalar.dtype, np.number):
                        record[key] = float(scalar.item())

                if isinstance(group, h5py.Group):
                    for key, item in group.items():
                        if isinstance(item, h5py.Dataset) and item.shape in [(), (1,)]:
                            value = np.asarray(item[()])
                            if value.size == 1 and np.issubdtype(value.dtype, np.number) and not np.issubdtype(value.dtype, np.bool_):
                                record[key] = float(value.item())

                shape = record.get("shape_score")
                if shape is None and "area" in record and "perimeter" in record and record["perimeter"] > 0:
                    shape = 4.0 * math.pi * record["area"] / record["perimeter"] ** 2
                    record["shape_score"] = shape

                convex_shape = record.get("convex_shape")
                if shape is not None and convex_shape is not None and convex_shape != 0:
                    shape_val = float(shape)
                    convex_val = float(convex_shape)
                    irregularity_val = math.atan(shape_val / convex_val)

                    record.update({
                        "Shape": shape_val,
                        "Convex Shape": convex_val,
                        "Irregularity": irregularity_val,
                    })
                    record.setdefault("irregularity", irregularity_val)
                    data.append(record)

    except Exception as e:
        print(f"Error al leer {os.path.basename(file_path)}: {str(e)}")
        return None

    if not data:
        return None

    df = pd.DataFrame(data)

    # Ordenar por Shape, Convex Shape e Irregularity
    df = df.sort_values(by=["Shape", "Convex Shape", "Irregularity"]).reset_index(drop=True)
    return df


def run_multivariate_analysis(df, sample_col="Muestra"):
    """Compare every sample pair using filtered, canonical numeric pore features."""
    excluded_cols = {
        "index",
        "ellipse_major_diameter_um",
        "rectangle_minor_side_um",
        "ellipse_minor_diameter_um",
        "equivalent_diameter_um",
        "rectangle_major_side_um",
        "pore_elongation",
        "contour_id",
        "ruta_archivo",
    }
    alias_groups = {
        "irregularity": ("irregularity", "Irregularity", "pore_irregularity_deg"),
        "convex_shape": ("convex_shape", "Convex Shape"),
        "shape_score": ("shape_score", "Shape"),
    }
    alias_names = {name for aliases in alias_groups.values() for name in aliases}
    numeric_cols = [
        column
        for column in df.columns
        if column != sample_col
        and column.casefold() not in excluded_cols
        and pd.api.types.is_numeric_dtype(df[column])
    ]

    features = pd.DataFrame(index=df.index)
    for canonical_name, aliases in alias_groups.items():
        selected = next(
            (
                alias
                for alias in aliases
                if alias in numeric_cols and df[alias].notna().any()
            ),
            None,
        )
        if selected is not None:
            values = df[selected]
            if selected == "pore_irregularity_deg":
                values = np.deg2rad(values)
            features[canonical_name] = values

    for column in numeric_cols:
        if column not in alias_names:
            features[column] = df[column]

    features = features.replace([np.inf, -np.inf], np.nan)
    feature_cols = list(features.columns)
    samples = df[sample_col].dropna().unique()
    if len(samples) < 2:
        raise ValueError("Se requieren al menos 2 muestras para comparar multivariadamente.")
    if not feature_cols:
        raise ValueError("No se encontraron variables numéricas para el análisis multivariado.")

    summary_records = []
    ranking_records = []
    for sample1_index, sample1_name in enumerate(samples[:-1]):
        for sample2_name in samples[sample1_index + 1:]:
            sample1_mask = df[sample_col] == sample1_name
            sample2_mask = df[sample_col] == sample2_name
            sample1 = features.loc[sample1_mask].dropna(subset=feature_cols)
            sample2 = features.loc[sample2_mask].dropna(subset=feature_cols)
            values1 = sample1[feature_cols].to_numpy(dtype=float)
            values2 = sample2[feature_cols].to_numpy(dtype=float)
            n1, n2 = len(values1), len(values2)
            feature_count = len(feature_cols)

            summary = {
                "Muestra 1": sample1_name,
                "Muestra 2": sample2_name,
                "N1 (is_edge=False)": n1,
                "N2 (is_edge=False)": n2,
                "Variables Analizadas (p)": feature_count,
                "Dist. Mahalanobis (D_M)": np.nan,
                "Hotelling T^2": np.nan,
                "F-Estadístico": np.nan,
                "p-valor (Hotelling)": np.nan,
                "Diferencia Multivariada": "No calculable (n < 2)",
                "RF AUC-ROC (OOB Score)": np.nan,
            }
            summary_records.append(summary)
            if n1 < 2 or n2 < 2:
                continue

            mean1, mean2 = np.mean(values1, axis=0), np.mean(values2, axis=0)
            difference = mean1 - mean2
            covariance1 = np.cov(values1, rowvar=False, ddof=1)
            covariance2 = np.cov(values2, rowvar=False, ddof=1)
            pooled_covariance = (
                (n1 - 1) * np.atleast_2d(covariance1) + (n2 - 1) * np.atleast_2d(covariance2)
            ) / (n1 + n2 - 2)
            mahalanobis_squared = float(difference @ np.linalg.pinv(pooled_covariance) @ difference)
            mahalanobis_distance = float(np.sqrt(max(0.0, mahalanobis_squared)))
            t_squared = (n1 * n2 / (n1 + n2)) * mahalanobis_squared

            df1 = feature_count
            df2 = n1 + n2 - feature_count - 1
            f_statistic = (
                (df2 / (feature_count * (n1 + n2 - 2))) * t_squared if df2 > 0 else np.nan
            )
            p_value = float(stats.f.sf(f_statistic, df1, df2)) if df2 > 0 else np.nan

            values = np.vstack([values1, values2])
            labels = np.array([0] * n1 + [1] * n2)
            classifier = RandomForestClassifier(n_estimators=200, random_state=42, oob_score=True)
            classifier.fit(values, labels)
            oob_predictions = classifier.oob_decision_function_[:, 1]
            valid_predictions = np.isfinite(oob_predictions)
            auc_score = (
                float(roc_auc_score(labels[valid_predictions], oob_predictions[valid_predictions]))
                if valid_predictions.sum() > 1 and len(np.unique(labels[valid_predictions])) == 2
                else np.nan
            )

            summary.update({
                "Dist. Mahalanobis (D_M)": mahalanobis_distance,
                "Hotelling T^2": t_squared,
                "F-Estadístico": f_statistic,
                "p-valor (Hotelling)": p_value,
                "Diferencia Multivariada": (
                    "Sí (p < 0.05)" if p_value < 0.05 else "No Significativa"
                ),
                "RF AUC-ROC (OOB Score)": auc_score,
            })

            pair_ranking = []
            for index, column in enumerate(feature_cols):
                column_mean1, column_std1 = sample1[column].mean(), sample1[column].std(ddof=1)
                column_mean2, column_std2 = sample2[column].mean(), sample2[column].std(ddof=1)
                pooled_std = np.sqrt(
                    ((n1 - 1) * column_std1**2 + (n2 - 1) * column_std2**2)
                    / (n1 + n2 - 2)
                )
                cohen_d = abs(column_mean1 - column_mean2) / pooled_std if pooled_std > 0 else 0.0
                pair_ranking.append({
                    "Muestra 1": sample1_name,
                    "Muestra 2": sample2_name,
                    "Ranking": None,
                    "Variable": column,
                    "Importancia RF (%)": classifier.feature_importances_[index] * 100,
                    "Tamaño Efecto (d de Cohen)": cohen_d,
                    f"Media ({sample1_name})": column_mean1,
                    f"DE ({sample1_name})": column_std1,
                    f"Media ({sample2_name})": column_mean2,
                    f"DE ({sample2_name})": column_std2,
                })
            pair_ranking.sort(key=lambda record: record["Importancia RF (%)"], reverse=True)
            for rank, record in enumerate(pair_ranking, start=1):
                record["Ranking"] = rank
                ranking_records.append(record)

    if not summary_records:
        raise ValueError("No se encontraron pares de muestras para comparar.")
    ranking_columns = [
        "Muestra 1",
        "Muestra 2",
        "Ranking",
        "Variable",
        "Importancia RF (%)",
        "Tamaño Efecto (d de Cohen)",
    ]
    ranking_df = pd.DataFrame(ranking_records)
    if not ranking_df.empty:
        ranking_columns.extend(column for column in ranking_df.columns if column not in ranking_columns)
        ranking_df = ranking_df[ranking_columns]
    return pd.DataFrame(summary_records), ranking_df


def add_multivariate_tab_to_excel(excel_path, summary, ranking, sheet_name="Análisis Multivariado Global"):
    """Add the multivariate summary and feature ranking to an existing workbook."""
    workbook = openpyxl.load_workbook(excel_path)
    if sheet_name in workbook.sheetnames:
        workbook.remove(workbook[sheet_name])
    worksheet = workbook.create_sheet(title=sheet_name)

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    title_font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    section_font = Font(name="Calibri", size=11, bold=True, color="1F4E79")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=10, bold=True)
    regular_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    worksheet.cell(row=1, column=1, value="ANÁLISIS ESTADÍSTICO MULTIVARIADO (p VARIABLES)").font = title_font
    worksheet.cell(row=2, column=1, value="Filtro aplicado: is_edge = False (poros interiores)").font = Font(
        italic=True, size=10, color="595959"
    )
    worksheet.cell(row=4, column=1, value="1. Resumen de Pruebas Multivariadas Globales").font = section_font

    summary_df = summary if isinstance(summary, pd.DataFrame) else pd.DataFrame([summary])
    for row_index, row in enumerate(
        dataframe_to_rows(summary_df, index=False, header=True), start=5
    ):
        worksheet.append(row)
        for cell in worksheet[worksheet.max_row]:
            cell.border = thin_border
            if row_index == 5:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.font = regular_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if isinstance(cell.value, float):
                    cell.value = round(cell.value, 6) if abs(cell.value) < 0.001 else round(cell.value, 4)

    ranking_start = worksheet.max_row + 3
    worksheet.cell(
        row=ranking_start, column=1,
        value="2. Ranking de Contribución por Variable (Random Forest + Cohen's d)",
    ).font = section_font
    header_row = ranking_start + 1
    for row_index, row in enumerate(
        dataframe_to_rows(ranking, index=False, header=True), start=header_row
    ):
        worksheet.append(row)
        for column_index, cell in enumerate(worksheet[worksheet.max_row], start=1):
            cell.border = thin_border
            if row_index == header_row:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.font = bold_font if column_index <= 3 else regular_font
                cell.alignment = Alignment(
                    horizontal=(
                        "left" if column_index in (1, 2, 4) else
                        "center" if column_index == 3 else "right"
                    ),
                    vertical="center",
                )
                if isinstance(cell.value, float):
                    cell.value = round(cell.value, 4)

    for column in worksheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column)
        column_letter = openpyxl.utils.get_column_letter(column[0].column)
        worksheet.column_dimensions[column_letter].width = max(max_length + 3, 14)
    workbook.save(excel_path)


# ==========================================
# 2. PRUEBAS ESTADÍSTICAS BIVARIADAS (2D)
# ==========================================
def ks2d2s(x1, y1, x2, y2):
    """
    Prueba Kolmogorov-Smirnov bidimensional (2D K-S) para comparar dos distribuciones
    bivariadas (Convex Shape vs. Irregularity). Devuelve el estadístico D y el p-valor.
    """
    x1, y1 = np.asarray(x1, dtype=float), np.asarray(y1, dtype=float)
    x2, y2 = np.asarray(x2, dtype=float), np.asarray(y2, dtype=float)
    
    n1, n2 = len(x1), len(x2)
    if n1 == 0 or n2 == 0:
        return np.nan, np.nan

    d1 = np.column_stack([x1, y1])
    d2 = np.column_stack([x2, y2])

    all_x = np.concatenate([x1, x2])
    all_y = np.concatenate([y1, y2])
    
    # Muestreo para acelerar el cálculo si hay muchos datos
    eval_x, eval_y = all_x, all_y
    if len(eval_x) > 400:
        idx = np.random.choice(len(eval_x), 400, replace=False)
        eval_x, eval_y = eval_x[idx], eval_y[idx]

    # Proporción de puntos en los 4 cuadrantes respecto a cada punto de evaluación
    f1_q1 = np.mean((d1[:, 0][:, None] <= eval_x[None, :]) & (d1[:, 1][:, None] <= eval_y[None, :]), axis=0)
    f2_q1 = np.mean((d2[:, 0][:, None] <= eval_x[None, :]) & (d2[:, 1][:, None] <= eval_y[None, :]), axis=0)

    f1_q2 = np.mean((d1[:, 0][:, None] <= eval_x[None, :]) & (d1[:, 1][:, None] >= eval_y[None, :]), axis=0)
    f2_q2 = np.mean((d2[:, 0][:, None] <= eval_x[None, :]) & (d2[:, 1][:, None] >= eval_y[None, :]), axis=0)

    f1_q3 = np.mean((d1[:, 0][:, None] >= eval_x[None, :]) & (d1[:, 1][:, None] <= eval_y[None, :]), axis=0)
    f2_q3 = np.mean((d2[:, 0][:, None] >= eval_x[None, :]) & (d2[:, 1][:, None] <= eval_y[None, :]), axis=0)

    f1_q4 = np.mean((d1[:, 0][:, None] >= eval_x[None, :]) & (d1[:, 1][:, None] >= eval_y[None, :]), axis=0)
    f2_q4 = np.mean((d2[:, 0][:, None] >= eval_x[None, :]) & (d2[:, 1][:, None] >= eval_y[None, :]), axis=0)

    d_max = max(
        np.max(np.abs(f1_q1 - f2_q1)),
        np.max(np.abs(f1_q2 - f2_q2)),
        np.max(np.abs(f1_q3 - f2_q3)),
        np.max(np.abs(f1_q4 - f2_q4))
    )

    # Tamaño muestral efectivo y aproximación de p-valor
    n_eff = (n1 * n2) / (n1 + n2)
    r1 = np.corrcoef(x1, y1)[0, 1] if n1 > 1 else 0
    r2 = np.corrcoef(x2, y2)[0, 1] if n2 > 1 else 0
    r = np.nanmean([r1, r2])
    if np.isnan(r):
        r = 0.0

    denom = 1.0 + np.sqrt(1.0 - r**2) * (0.25 - 0.75 / np.sqrt(n_eff))
    z = d_max * np.sqrt(n_eff) / denom if denom > 0 else d_max * np.sqrt(n_eff)
    p_val = 2.0 * np.exp(-2.0 * (z ** 2))
    p_val = np.clip(p_val, 0.0, 1.0)

    return d_max, p_val


def hotelling_t2_2samp(x1, y1, x2, y2):
    """
    Prueba T^2 de Hotelling para dos muestras independientes bivariadas.
    Retorna: Distancia de Mahalanobis (D_M), Estadístico T^2, Estadístico F y p-valor.
    """
    X1 = np.column_stack([x1, y1])
    X2 = np.column_stack([x2, y2])
    
    n1, p = X1.shape
    n2, _ = X2.shape
    
    if n1 <= p or n2 <= p:
        return np.nan, np.nan, np.nan, np.nan
        
    mean1 = np.mean(X1, axis=0)
    mean2 = np.mean(X2, axis=0)
    diff = mean1 - mean2
    
    cov1 = np.cov(X1, rowvar=False, ddof=1)
    cov2 = np.cov(X2, rowvar=False, ddof=1)
    
    # Matriz de covarianza combinada (pooled covariance)
    S_p = ((n1 - 1) * cov1 + (n2 - 1) * cov2) / (n1 + n2 - 2)
    
    try:
        S_p_inv = np.linalg.pinv(S_p)
    except Exception:
        return np.nan, np.nan, np.nan, np.nan
        
    d_m_sq = float(diff.T @ S_p_inv @ diff)
    d_m = np.sqrt(max(0.0, d_m_sq))
    
    t2 = (n1 * n2 / (n1 + n2)) * d_m_sq
    
    df1 = p
    df2 = n1 + n2 - p - 1
    
    if df2 > 0:
        f_stat = ((n1 + n2 - p - 1) / (p * (n1 + n2 - 2))) * t2
        p_val = stats.f.sf(f_stat, df1, df2)
    else:
        f_stat = np.nan
        p_val = np.nan
        
    return d_m, t2, f_stat, p_val


# ==========================================
# 3. PIPELINE PRINCIPAL Y EXPORTACIÓN COMPLETA
# ==========================================
def process_all_mosaics(root_directory, output_excel="Resultados_Poros_Global.xlsx"):
    root_directory = os.path.abspath(os.path.expanduser(root_directory))

    if not os.path.exists(root_directory):
        print(f"❌ ERROR: La ruta especificada NO existe o no se puede acceder:")
        print(f"   '{root_directory}'")
        return

    print(f"Buscando archivos .h5 / .hdf5 en: {root_directory}\n")

    h5_files = []
    for dirpath, _, filenames in os.walk(root_directory):
        for filename in filenames:
            if filename.lower().endswith((".h5", ".hdf5")):
                h5_files.append(os.path.join(dirpath, filename))

    if not h5_files:
        print("⚠️ No se encontraron archivos .h5 o .hdf5 en la carpeta.")
        return

    print(f"✅ Se encontraron {len(h5_files)} archivos de datos. Procesando...")

    file_dfs = {}
    for file_path in h5_files:
        filename_only = os.path.basename(file_path)
        df = process_h5_file(file_path)
        
        if df is not None and not df.empty:
            file_dfs[filename_only] = df
            print(f"  ✓ {filename_only}: {len(df)} poros interiores procesados.")
        else:
            print(f"  ✗ {filename_only}: Sin datos procesables.")

    if not file_dfs:
        print("\nNo se pudo extraer información de ningún archivo.")
        return

    # ------------------------------------------
    # A. Planilla de Datos Poros (Raw Data)
    # ------------------------------------------
    combined_columns = []
    for filename, df in file_dfs.items():
        df_copy = df.copy()
        df_copy.columns = pd.MultiIndex.from_tuples([(filename, col) for col in df_copy.columns])
        combined_columns.append(df_copy)

    final_excel_df = pd.concat(combined_columns, axis=1)

    # ------------------------------------------
    # B. Estadísticas Descriptivas por Muestra
    # ------------------------------------------
    summary_rows = []
    for filename, df in file_dfs.items():
        clean_name = filename.replace(".h5", "").replace(".hdf5", "")
        row = {
            'Muestra': clean_name,
            'Cantidad Poros (n)': len(df)
        }
        for col in ['Shape', 'Convex Shape', 'Irregularity']:
            vals = df[col]
            row[f'{col} Promedio'] = round(vals.mean(), 4)
            row[f'{col} Desv.Std'] = round(vals.std(), 4)
            row[f'{col} Mediana'] = round(vals.median(), 4)
            row[f'{col} Q25'] = round(vals.quantile(0.25), 4)
            row[f'{col} Q75'] = round(vals.quantile(0.75), 4)
        summary_rows.append(row)

    df_summary = pd.DataFrame(summary_rows)

    # ------------------------------------------
    # C. Pruebas Comparativas 2D (K-S 2D y Hotelling T^2)
    # ------------------------------------------
    print("\nCalculando pruebas bivariadas K-S 2D y Hotelling T^2 (todos contra todos)...")
    ks_rows = []
    filenames_list = list(file_dfs.keys())

    for i in range(len(filenames_list)):
        for j in range(i + 1, len(filenames_list)):
            f1, f2 = filenames_list[i], filenames_list[j]
            d1, d2 = file_dfs[f1], file_dfs[f2]
            
            # 1. Prueba K-S 2D (No Paramétrica)
            d_stat, p_val_ks = ks2d2s(
                d1["Convex Shape"], d1["Irregularity"],
                d2["Convex Shape"], d2["Irregularity"]
            )
            
            # 2. Prueba T^2 de Hotelling & Distancia de Mahalanobis (Paramétrica)
            d_m, t2_stat, f_stat, p_val_t2 = hotelling_t2_2samp(
                d1["Convex Shape"], d1["Irregularity"],
                d2["Convex Shape"], d2["Irregularity"]
            )
            
            # Formateo legible de p-valores
            p_ks_str = f"{p_val_ks:.2e}" if (not np.isnan(p_val_ks) and p_val_ks < 0.001) else (round(p_val_ks, 4) if not np.isnan(p_val_ks) else "N/A")
            p_t2_str = f"{p_val_t2:.2e}" if (not np.isnan(p_val_t2) and p_val_t2 < 0.001) else (round(p_val_t2, 4) if not np.isnan(p_val_t2) else "N/A")
            
            ks_rows.append({
                "Muestra 1": f1.replace(".h5", "").replace(".hdf5", ""),
                "Muestra 2": f2.replace(".h5", "").replace(".hdf5", ""),
                # Indicadores Kolmogorov-Smirnov 2D
                "Estadístico D (K-S 2D)": round(d_stat, 4) if not np.isnan(d_stat) else "N/A",
                "p-valor (K-S 2D)": p_ks_str,
                "Sig. K-S 2D (α=0.05)": "Sí" if (not np.isnan(p_val_ks) and p_val_ks < 0.05) else "No",
                # Indicadores Hotelling T^2 y Mahalanobis
                "Dist. Mahalanobis (D_M)": round(d_m, 4) if not np.isnan(d_m) else "N/A",
                "Estadístico T² (Hotelling)": round(t2_stat, 2) if not np.isnan(t2_stat) else "N/A",
                "p-valor (Hotelling T²)": p_t2_str,
                "Sig. Hotelling T² (α=0.05)": "Sí" if (not np.isnan(p_val_t2) and p_val_t2 < 0.05) else "No"
            })

    df_ks = pd.DataFrame(ks_rows)

    # ------------------------------------------
    # D. Generación de Gráficos (Archivos Temporales)
    # ------------------------------------------
    img_scatter_path = os.path.join(root_directory, "temp_scatter.png")
    img_kde_path = os.path.join(root_directory, "temp_kde.png")

    # 1. Gráfico de Dispersión
    plt.figure(figsize=(10, 6))
    for filename, df in file_dfs.items():
        plt.scatter(
            df["Convex Shape"], 
            df["Irregularity"], 
            label=filename.replace(".h5", "").replace(".hdf5", ""), 
            alpha=0.6, 
            edgecolors='none',
            s=25
        )
    plt.xlabel("Convex Shape", fontsize=12)
    plt.ylabel("Irregularity = atan(Shape / Convex Shape)", fontsize=12)
    plt.title("Irregularity vs. Convex Shape por Archivo", fontsize=14, fontweight='bold')
    plt.grid(True, linestyle="--", alpha=0.5)
    plt.legend(title="Archivos", bbox_to_anchor=(1.05, 1), loc='upper left')
    plt.tight_layout()
    plt.savefig(img_scatter_path, dpi=200, bbox_inches='tight')
    plt.close()

    # 2. Mosaico de Mapas de Densidad Kernel 2D (KDE 2D)
    num_samples = len(file_dfs)
    cols = 2 if num_samples > 1 else 1
    rows = math.ceil(num_samples / cols)

    fig, axes = plt.subplots(rows, cols, figsize=(6 * cols, 5 * rows), sharex=True, sharey=True)
    if num_samples == 1:
        axes = np.array([axes])
    axes_flat = axes.flatten()

    for idx, (filename, df) in enumerate(file_dfs.items()):
        ax = axes_flat[idx]
        clean_name = filename.replace(".h5", "").replace(".hdf5", "")
        
        sns.kdeplot(
            data=df, x="Convex Shape", y="Irregularity",
            ax=ax, cmap="Blues", fill=True, thresh=0.05, levels=10
        )
        ax.scatter(df["Convex Shape"], df["Irregularity"], s=6, alpha=0.15, color='darkblue')
        ax.set_title(f"Muestra: {clean_name} (n={len(df)})", fontsize=11, fontweight='bold')
        ax.set_xlabel("Convex Shape", fontsize=10)
        ax.set_ylabel("Irregularity = atan(Shape / Convex Shape)", fontsize=10)
        ax.grid(True, linestyle="--", alpha=0.4)
        ax.set_xlim(0, 1)
        ax.set_ylim(0, 0.85)

    for idx in range(num_samples, len(axes_flat)):
        axes_flat[idx].axis('off')

    plt.suptitle("Mapas de Densidad Kernel 2D (Irregularity vs. Convex Shape)", fontsize=14, fontweight='bold', y=0.99)
    plt.tight_layout()
    plt.savefig(img_kde_path, dpi=200, bbox_inches='tight')
    plt.close()

    # ------------------------------------------
    # E. Exportación a Excel con openpyxl e Inserción de Imágenes
    # ------------------------------------------
    excel_path = os.path.join(root_directory, output_excel)

    with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
        final_excel_df.to_excel(writer, sheet_name="Datos Poros")
        df_summary.to_excel(writer, sheet_name="Estadísticas Descriptivas", index=False)
        df_ks.to_excel(writer, sheet_name="Prueba K-S 2D", index=False)

    # Reabrir el libro para incrustar las imágenes en la pestaña de Gráficos
    wb = openpyxl.load_workbook(excel_path)
    ws_chart = wb.create_sheet(title="Gráficos")

    # Incrustar Gráfico 1: Dispersión
    ws_chart.cell(row=2, column=2, value="1. Gráfico de Dispersión Global (Irregularity vs. Convex Shape)")
    ws_chart.cell(row=2, column=2).font = openpyxl.styles.Font(bold=True, size=12)
    img_scatter = Image(img_scatter_path)
    ws_chart.add_image(img_scatter, "B4")

    # Incrustar Gráfico 2: Mosaico KDE 2D
    ws_chart.cell(row=38, column=2, value="2. Mapas de Densidad Kernel 2D (KDE 2D) - Mosaico Comparativo")
    ws_chart.cell(row=38, column=2).font = openpyxl.styles.Font(bold=True, size=12)
    img_kde = Image(img_kde_path)
    ws_chart.add_image(img_kde, "B40")

    wb.save(excel_path)

    if len(file_dfs) >= 2:
        multivariate_df = pd.concat(
            [
                df.assign(Muestra=os.path.splitext(filename)[0])
                for filename, df in file_dfs.items()
            ],
            ignore_index=True,
        )
        try:
            summary, ranking = run_multivariate_analysis(multivariate_df)
            add_multivariate_tab_to_excel(excel_path, summary, ranking)
        except ValueError as error:
            print(f"No se pudo generar el análisis multivariado: {error}")
    else:
        print("Análisis multivariado omitido: se necesitan al menos 2 muestras.")

    # Limpiar archivos de imagen temporales
    if os.path.exists(img_scatter_path):
        os.remove(img_scatter_path)
    if os.path.exists(img_kde_path):
        os.remove(img_kde_path)

    print(f"\n¡Proceso completado exitosamente!")
    print(f"📊 Archivo Excel generado: {excel_path}")


# ==========================================
# PUNTO DE EJECUCIÓN
# ==========================================
if __name__ == "__main__":
    CARPETA_RAIZ = r"H:\Otros ordenadores\Microscopio\Silvia EEUU\Porosidad de los mosaicos 4"
    process_all_mosaics(CARPETA_RAIZ)