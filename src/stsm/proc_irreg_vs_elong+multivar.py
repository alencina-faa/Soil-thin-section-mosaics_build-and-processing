import math
import os
import h5py
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd
import seaborn as sns
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from scipy import stats
from sklearn.ensemble import RandomForestClassifier
from sklearn.inspection import permutation_importance
from sklearn.metrics import roc_auc_score


# ==========================================
# 1. FUNCIÓN DE LECTURA Y EXTRACCIÓN .H5
# ==========================================
def process_h5_file(file_path):
    """
    Lee un archivo .h5, extrae las variables de los poros interiores (is_edge==False),
    y calcula Shape, Convex Shape e Irregularity = atan(Shape / Convex Shape).
    """
    data = []

    try:
        with h5py.File(file_path, "r") as f:
            if "contours" not in f:
                return None

            contours_group = f["contours"]

            for contour_id, group in contours_group.items():
                # 1. Filtrar únicamente contornos interiores (is_edge = False)
                is_edge_value = np.asarray(group.attrs.get("is_edge", True))
                is_edge = bool(is_edge_value.reshape(-1)[0]) if is_edge_value.size else False
                if is_edge:
                    continue

                is_over_50um_value = np.asarray(group.attrs.get("is_over_50um", False))
                is_over_50um = (
                    bool(is_over_50um_value.reshape(-1)[0]) if is_over_50um_value.size else False
                )
                record = {
                    "Contour_ID": contour_id,
                    "is_edge": is_edge,
                    "is_over_50um": is_over_50um,
                }
                for key, value in group.attrs.items():
                    if key in {"is_edge", "is_over_50um"}:
                        continue
                    scalar = np.asarray(value)
                    if scalar.size == 1 and np.issubdtype(scalar.dtype, np.number):
                        record[key] = float(scalar.item())

                if isinstance(group, h5py.Group):
                    for key, item in group.items():
                        if isinstance(item, h5py.Dataset) and item.shape in [(), (1,)]:
                            value = np.asarray(item[()])
                            if value.size == 1 and np.issubdtype(value.dtype, np.number) and not np.issubdtype(value.dtype, np.bool_):
                                record[key] = float(value.item())

                shape = record.get("shape_score")
                if shape is None and "area" in record and "perimeter" in record and record["perimeter"] > 0:
                    shape = 4.0 * math.pi * record["area"] / record["perimeter"] ** 2
                    record["shape_score"] = shape

                convex_shape = record.get("convex_shape")
                if shape is not None and convex_shape is not None and convex_shape != 0:
                    shape_val = float(shape)
                    convex_val = float(convex_shape)
                    irregularity_val = math.atan(shape_val / convex_val)

                    record.update({
                        "Shape": shape_val,
                        "Convex Shape": convex_val,
                        "Irregularity": irregularity_val,
                    })
                    record.setdefault("irregularity", irregularity_val)
                    data.append(record)

    except Exception as e:
        print(f"Error al leer {os.path.basename(file_path)}: {str(e)}")
        return None

    if not data:
        return None

    df = pd.DataFrame(data)

    # Ordenar por Shape, Convex Shape e Irregularity
    df = df.sort_values(by=["Shape", "Convex Shape", "Irregularity"]).reset_index(drop=True)
    return df


# ==========================================
# 2. ANÁLISIS MULTIVARIADO (HOTELLING T^2 + PERMUTATION IMPORTANCE)
# ==========================================
def run_multivariate_analysis(df, sample_col="Muestra"):
    """
    Compara todos los pares de muestras en p dimensiones usando T^2 de Hotelling
    e Importancia por Permutación (Permutation Feature Importance) basada en ROC-AUC.
    """
    candidate_features = [
        "area",
        "convex_shape",
        "ellipse_angle_deg",
        "perimeter",
        "pore_elongation",
        "pore_irregularity_deg",
        "shape_score",
    ]
    feature_cols = [c for c in candidate_features if c in df.columns]
    if not feature_cols:
        ignore_cols = {sample_col, "Contour_ID", "is_edge", "is_over_50um"}
        feature_cols = sorted([c for c in df.columns if c not in ignore_cols and pd.api.types.is_numeric_dtype(df[c])])

    eligible_mask = pd.Series(True, index=df.index)
    if "is_edge" in df.columns:
        eligible_mask &= df["is_edge"].eq(False)
    if "is_over_50um" in df.columns:
        eligible_mask &= df["is_over_50um"].eq(True)

    eligible_df = df.loc[eligible_mask].copy()
    features = eligible_df[feature_cols].replace([np.inf, -np.inf], np.nan)
    samples = eligible_df[sample_col].dropna().unique()

    if len(samples) < 2:
        raise ValueError(
            "Se requieren al menos 2 muestras con is_edge=False e is_over_50um=True para el análisis multivariado."
        )

    summary_records = []
    ranking_records = []
    detail_records = []

    for sample1_index, sample1_name in enumerate(samples[:-1]):
        for sample2_name in samples[sample1_index + 1:]:
            sample1_mask = eligible_df[sample_col] == sample1_name
            sample2_mask = eligible_df[sample_col] == sample2_name
            sample1 = features.loc[sample1_mask].dropna(subset=feature_cols)
            sample2 = features.loc[sample2_mask].dropna(subset=feature_cols)
            values1 = sample1[feature_cols].to_numpy(dtype=float)
            values2 = sample2[feature_cols].to_numpy(dtype=float)
            n1, n2 = len(values1), len(values2)
            feature_count = len(feature_cols)

            summary = {
                "Muestra 1": sample1_name,
                "Muestra 2": sample2_name,
                "N1 (condiciones aplicadas)": n1,
                "N2 (condiciones aplicadas)": n2,
                "Variables Analizadas (p)": feature_count,
                "Dist. Mahalanobis (D_M)": np.nan,
                "Hotelling T^2": np.nan,
                "F-Estadístico": np.nan,
                "p-valor (Hotelling)": np.nan,
                "Diferencia Multivariada": "No calculable (n < 2)",
                "RF AUC-ROC (OOB Score)": np.nan,
            }
            summary_records.append(summary)

            if n1 < 2 or n2 < 2:
                continue

            # 1. Prueba T^2 de Hotelling & Mahalanobis
            mean1, mean2 = np.mean(values1, axis=0), np.mean(values2, axis=0)
            difference = mean1 - mean2
            covariance1 = np.cov(values1, rowvar=False, ddof=1)
            covariance2 = np.cov(values2, rowvar=False, ddof=1)
            pooled_covariance = (
                (n1 - 1) * np.atleast_2d(covariance1) + (n2 - 1) * np.atleast_2d(covariance2)
            ) / (n1 + n2 - 2)
            mahalanobis_squared = float(difference @ np.linalg.pinv(pooled_covariance) @ difference)
            mahalanobis_distance = float(np.sqrt(max(0.0, mahalanobis_squared)))
            t_squared = (n1 * n2 / (n1 + n2)) * mahalanobis_squared

            df1 = feature_count
            df2 = n1 + n2 - feature_count - 1
            f_statistic = (
                (df2 / (feature_count * (n1 + n2 - 2))) * t_squared if df2 > 0 else np.nan
            )
            p_value = float(stats.f.sf(f_statistic, df1, df2)) if df2 > 0 else np.nan

            # 2. Random Forest Classifier
            values = np.vstack([values1, values2])
            labels = np.array([0] * n1 + [1] * n2)
            classifier = RandomForestClassifier(n_estimators=200, random_state=42, oob_score=True)
            classifier.fit(values, labels)

            oob_predictions = classifier.oob_decision_function_[:, 1]
            valid_predictions = np.isfinite(oob_predictions)
            auc_score = (
                float(roc_auc_score(labels[valid_predictions], oob_predictions[valid_predictions]))
                if valid_predictions.sum() > 1 and len(np.unique(labels[valid_predictions])) == 2
                else np.nan
            )

            # 3. Permutation Feature Importance (Basada en la caída de ROC-AUC)
            perm_res = permutation_importance(
                classifier, values, labels,
                scoring='roc_auc',
                n_repeats=10,
                random_state=42
            )
            drop_auc_mean = perm_res.importances_mean
            drop_auc_clipped = np.maximum(0, drop_auc_mean)
            if drop_auc_clipped.sum() > 0:
                importances_pct = (drop_auc_clipped / drop_auc_clipped.sum()) * 100.0
            else:
                importances_pct = np.zeros_like(drop_auc_clipped)

            summary.update({
                "Dist. Mahalanobis (D_M)": mahalanobis_distance,
                "Hotelling T^2": t_squared,
                "F-Estadístico": f_statistic,
                "p-valor (Hotelling)": p_value,
                "Diferencia Multivariada": (
                    "Sí (p < 0.05)" if p_value < 0.05 else "No Significativa"
                ),
                "RF AUC-ROC (OOB Score)": auc_score,
            })

            pair_records = []
            for index, column in enumerate(feature_cols):
                column_mean1, column_std1 = sample1[column].mean(), sample1[column].std(ddof=1)
                column_mean2, column_std2 = sample2[column].mean(), sample2[column].std(ddof=1)
                pooled_std = np.sqrt(
                    ((n1 - 1) * column_std1**2 + (n2 - 1) * column_std2**2) / (n1 + n2 - 2)
                )
                cohen_d = abs(column_mean1 - column_mean2) / pooled_std if pooled_std > 0 else 0.0

                # Registro para la Matriz Pivoteada
                ranking_records.append({
                    "Muestra 1": sample1_name,
                    "Muestra 2": sample2_name,
                    "Variable": column,
                    "Importancia Permutación RF (%)": importances_pct[index]
                })

                # Registro para la Tabla de Detalles Estadísticos
                pair_records.append({
                    "Muestra 1": sample1_name,
                    "Muestra 2": sample2_name,
                    "Variable": column,
                    "Tamaño Efecto (d de Cohen)": cohen_d,
                    f"Media ({sample1_name})": column_mean1,
                    f"DE ({sample1_name})": column_std1,
                    f"Media ({sample2_name})": column_mean2,
                    f"DE ({sample2_name})": column_std2,
                    "Importancia Permutación RF (%)": importances_pct[index]
                })

            pair_records.sort(key=lambda record: record["Importancia Permutación RF (%)"], reverse=True)
            for rank, record in enumerate(pair_records, start=1):
                del record["Importancia Permutación RF (%)"]
                record_with_rank = {"Ranking": rank}
                record_with_rank.update(record)
                detail_records.append(record_with_rank)

    # Crear Matriz Pivoteada (Variables en Filas, Pares Muestra 1 vs Muestra 2 en Columnas)
    ranking_df = pd.DataFrame(ranking_records)
    pivot_matrix_df = ranking_df.pivot(
        index="Variable", 
        columns=["Muestra 1", "Muestra 2"], 
        values="Importancia Permutación RF (%)"
    )

    return pd.DataFrame(summary_records), pivot_matrix_df, pd.DataFrame(detail_records)


def add_multivariate_tab_to_excel(excel_path, summary, ranking_pivot, detail, sheet_name="Análisis Multivariado Global"):
    """
    Añade la pestaña 'Análisis Multivariado Global' organizando el resumen,
    la matriz pivoteada de importancia de variables (estilo imagen) y los detalles univariados.
    """
    workbook = openpyxl.load_workbook(excel_path)
    if sheet_name in workbook.sheetnames:
        workbook.remove(workbook[sheet_name])
    worksheet = workbook.create_sheet(title=sheet_name)

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    title_font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    section_font = Font(name="Calibri", size=11, bold=True, color="1F4E79")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=10, bold=True)
    regular_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # Título principal
    worksheet.cell(row=1, column=1, value="ANÁLISIS ESTADÍSTICO MULTIVARIADO (p VARIABLES)").font = title_font
    worksheet.cell(
        row=2,
        column=1,
        value="Filtros aplicados: is_edge = False e is_over_50um = True | Importancia calculada por Permutación (roc_auc)",
    ).font = Font(italic=True, size=10, color="595959")

    # 1. Resumen Global
    worksheet.cell(row=4, column=1, value="1. Resumen de Pruebas Multivariadas Globales").font = section_font
    summary_df = summary if isinstance(summary, pd.DataFrame) else pd.DataFrame([summary])
    for row_index, row in enumerate(
        dataframe_to_rows(summary_df, index=False, header=True), start=5
    ):
        worksheet.append(row)
        for cell in worksheet[worksheet.max_row]:
            cell.border = thin_border
            if row_index == 5:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.font = regular_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if isinstance(cell.value, float):
                    cell.value = round(cell.value, 6) if abs(cell.value) < 0.001 else round(cell.value, 4)

    # 2. Matriz de Contribución por Variable (Pivoteada estilo Imagen)
    ranking_start = worksheet.max_row + 3
    worksheet.cell(
        row=ranking_start, column=1,
        value="2. Matriz de Contribución por Variable - Importancia Permutación RF (%)",
    ).font = section_font

    header_row_1 = ranking_start + 1
    header_row_2 = ranking_start + 2

    # Celda esquina superior izquierda (Variable)
    c_var1 = worksheet.cell(row=header_row_1, column=1, value="Variable")
    c_var1.fill = header_fill
    c_var1.font = header_font
    c_var1.alignment = Alignment(horizontal="center", vertical="center")
    c_var1.border = thin_border

    c_var2 = worksheet.cell(row=header_row_2, column=1, value="")
    c_var2.fill = header_fill
    c_var2.font = header_font
    c_var2.border = thin_border

    # Escribir pares en los encabezados de columnas
    pairs = ranking_pivot.columns  # MultiIndex (Muestra 1, Muestra 2)
    for col_idx, (m1, m2) in enumerate(pairs, start=2):
        cell_m1 = worksheet.cell(row=header_row_1, column=col_idx, value=str(m1))
        cell_m1.fill = header_fill
        cell_m1.font = header_font
        cell_m1.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_m1.border = thin_border

        cell_m2 = worksheet.cell(row=header_row_2, column=col_idx, value=str(m2))
        cell_m2.fill = header_fill
        cell_m2.font = header_font
        cell_m2.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_m2.border = thin_border

    # Escribir filas de variables y valores
    for row_idx, (var_name, row_series) in enumerate(ranking_pivot.iterrows(), start=header_row_2 + 1):
        c_var = worksheet.cell(row=row_idx, column=1, value=str(var_name))
        c_var.font = bold_font
        c_var.alignment = Alignment(horizontal="left", vertical="center")
        c_var.border = thin_border

        for col_idx, val in enumerate(row_series, start=2):
            c_val = worksheet.cell(row=row_idx, column=col_idx, value=float(val) if pd.notnull(val) else None)
            c_val.font = regular_font
            c_val.alignment = Alignment(horizontal="right", vertical="center")
            c_val.border = thin_border
            if c_val.value is not None:
                c_val.number_format = "0.0000"

    # 3. Detalles Estadísticos Univariados (d de Cohen, Medias, DE)
    detail_start = worksheet.max_row + 3
    worksheet.cell(
        row=detail_start, column=1,
        value="3. Detalles Estadísticos Univariados por Variable (Tamaño de Efecto d de Cohen, Medias y DE)",
    ).font = section_font
    header_row_detail = detail_start + 1
    for row_index, row in enumerate(
        dataframe_to_rows(detail, index=False, header=True), start=header_row_detail
    ):
        worksheet.append(row)
        for column_index, cell in enumerate(worksheet[worksheet.max_row], start=1):
            cell.border = thin_border
            if row_index == header_row_detail:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.font = bold_font if column_index <= 3 else regular_font
                cell.alignment = Alignment(
                    horizontal=(
                        "left" if column_index in (1, 2, 4) else
                        "center" if column_index == 3 else "right"
                    ),
                    vertical="center",
                )
                if isinstance(cell.value, float):
                    cell.value = round(cell.value, 4)

    # Ajustar ancho automático de columnas
    for column in worksheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column)
        column_letter = openpyxl.utils.get_column_letter(column[0].column)
        worksheet.column_dimensions[column_letter].width = max(max_length + 3, 14)

    workbook.save(excel_path)


# ==========================================
# 3. PRUEBAS ESTADÍSTICAS BIVARIADAS (2D)
# ==========================================
def ks2d2s(x1, y1, x2, y2):
    x1, y1 = np.asarray(x1, dtype=float), np.asarray(y1, dtype=float)
    x2, y2 = np.asarray(x2, dtype=float), np.asarray(y2, dtype=float)
    n1, n2 = len(x1), len(x2)
    if n1 == 0 or n2 == 0:
        return np.nan, np.nan
    d1 = np.column_stack([x1, y1])
    d2 = np.column_stack([x2, y2])
    all_x = np.concatenate([x1, x2])
    all_y = np.concatenate([y1, y2])
    eval_x, eval_y = all_x, all_y
    if len(eval_x) > 400:
        idx = np.random.choice(len(eval_x), 400, replace=False)
        eval_x, eval_y = eval_x[idx], eval_y[idx]
    f1_q1 = np.mean((d1[:, 0][:, None] <= eval_x[None, :]) & (d1[:, 1][:, None] <= eval_y[None, :]), axis=0)
    f2_q1 = np.mean((d2[:, 0][:, None] <= eval_x[None, :]) & (d2[:, 1][:, None] <= eval_y[None, :]), axis=0)
    f1_q2 = np.mean((d1[:, 0][:, None] <= eval_x[None, :]) & (d1[:, 1][:, None] >= eval_y[None, :]), axis=0)
    f2_q2 = np.mean((d2[:, 0][:, None] <= eval_x[None, :]) & (d2[:, 1][:, None] >= eval_y[None, :]), axis=0)
    f1_q3 = np.mean((d1[:, 0][:, None] >= eval_x[None, :]) & (d1[:, 1][:, None] <= eval_y[None, :]), axis=0)
    f2_q3 = np.mean((d2[:, 0][:, None] >= eval_x[None, :]) & (d2[:, 1][:, None] <= eval_y[None, :]), axis=0)
    f1_q4 = np.mean((d1[:, 0][:, None] >= eval_x[None, :]) & (d1[:, 1][:, None] >= eval_y[None, :]), axis=0)
    f2_q4 = np.mean((d2[:, 0][:, None] >= eval_x[None, :]) & (d2[:, 1][:, None] >= eval_y[None, :]), axis=0)
    d_max = max(
        np.max(np.abs(f1_q1 - f2_q1)),
        np.max(np.abs(f1_q2 - f2_q2)),
        np.max(np.abs(f1_q3 - f2_q3)),
        np.max(np.abs(f1_q4 - f2_q4))
    )
    n_eff = (n1 * n2) / (n1 + n2)
    r1 = np.corrcoef(x1, y1)[0, 1] if n1 > 1 else 0
    r2 = np.corrcoef(x2, y2)[0, 1] if n2 > 1 else 0
    r = np.nanmean([r1, r2])
    if np.isnan(r):
        r = 0.0
    denom = 1.0 + np.sqrt(1.0 - r**2) * (0.25 - 0.75 / np.sqrt(n_eff))
    z = d_max * np.sqrt(n_eff) / denom if denom > 0 else d_max * np.sqrt(n_eff)
    p_val = 2.0 * np.exp(-2.0 * (z ** 2))
    p_val = np.clip(p_val, 0.0, 1.0)
    return d_max, p_val


def hotelling_t2_2samp(x1, y1, x2, y2):
    X1 = np.column_stack([x1, y1])
    X2 = np.column_stack([x2, y2])
    n1, p = X1.shape
    n2, _ = X2.shape
    if n1 <= p or n2 <= p:
        return np.nan, np.nan, np.nan, np.nan
    mean1 = np.mean(X1, axis=0)
    mean2 = np.mean(X2, axis=0)
    diff = mean1 - mean2
    cov1 = np.cov(X1, rowvar=False, ddof=1)
    cov2 = np.cov(X2, rowvar=False, ddof=1)
    S_p = ((n1 - 1) * cov1 + (n2 - 1) * cov2) / (n1 + n2 - 2)
    try:
        S_p_inv = np.linalg.pinv(S_p)
    except Exception:
        return np.nan, np.nan, np.nan, np.nan
    d_m_sq = float(diff.T @ S_p_inv @ diff)
    d_m = np.sqrt(max(0.0, d_m_sq))
    t2 = (n1 * n2 / (n1 + n2)) * d_m_sq
    df1 = p
    df2 = n1 + n2 - p - 1
    if df2 > 0:
        f_stat = ((n1 + n2 - p - 1) / (p * (n1 + n2 - 2))) * t2
        p_val = stats.f.sf(f_stat, df1, df2)
    else:
        f_stat = np.nan
        p_val = np.nan
    return d_m, t2, f_stat, p_val


# ==========================================
# 4. PIPELINE PRINCIPAL Y EXPORTACIÓN COMPLETA
# ==========================================
def process_all_mosaics(root_directory, output_excel="Resultados_Poros_Global.xlsx"):
    root_directory = os.path.abspath(os.path.expanduser(root_directory))
    if not os.path.exists(root_directory):
        print(f"❌ ERROR: La ruta especificada NO existe: '{root_directory}'")
        return

    print(f"Buscando archivos .h5 / .hdf5 en: {root_directory}\n")
    h5_files = []
    for dirpath, _, filenames in os.walk(root_directory):
        for filename in filenames:
            if filename.lower().endswith((".h5", ".hdf5")):
                h5_files.append(os.path.join(dirpath, filename))

    if not h5_files:
        print("⚠️ No se encontraron archivos .h5 o .hdf5 en la carpeta.")
        return

    print(f"✅ Se encontraron {len(h5_files)} archivos de datos. Procesando...")
    file_dfs = {}
    for file_path in h5_files:
        filename_only = os.path.basename(file_path)
        df = process_h5_file(file_path)
        if df is not None and not df.empty:
            file_dfs[filename_only] = df
            print(f"  ✓ {filename_only}: {len(df)} poros interiores procesados.")
        else:
            print(f"  ✗ {filename_only}: Sin datos procesables.")

    if not file_dfs:
        print("\nNo se pudo extraer información de ningún archivo.")
        return

    # A. Planilla de Datos Poros (Raw Data)
    combined_columns = []
    for filename, df in file_dfs.items():
        df_copy = df.copy()
        df_copy.columns = pd.MultiIndex.from_tuples([(filename, col) for col in df_copy.columns])
        combined_columns.append(df_copy)
    final_excel_df = pd.concat(combined_columns, axis=1)

    # B. Estadísticas Descriptivas
    summary_rows = []
    for filename, df in file_dfs.items():
        clean_name = filename.replace(".h5", "").replace(".hdf5", "")
        row = {'Muestra': clean_name, 'Cantidad Poros (n)': len(df)}
        for col in ['Shape', 'Convex Shape', 'Irregularity']:
            if col in df.columns:
                vals = df[col]
                row[f'{col} Promedio'] = round(vals.mean(), 4)
                row[f'{col} Desv.Std'] = round(vals.std(), 4)
                row[f'{col} Mediana'] = round(vals.median(), 4)
                row[f'{col} Q25'] = round(vals.quantile(0.25), 4)
                row[f'{col} Q75'] = round(vals.quantile(0.75), 4)
        summary_rows.append(row)
    df_summary = pd.DataFrame(summary_rows)

    # C. Pruebas Comparativas 2D (K-S 2D y Hotelling T^2 2D)
    ks_rows = []
    filenames_list = list(file_dfs.keys())
    for i in range(len(filenames_list)):
        for j in range(i + 1, len(filenames_list)):
            f1, f2 = filenames_list[i], filenames_list[j]
            d1, d2 = file_dfs[f1], file_dfs[f2]
            d_stat, p_val_ks = ks2d2s(
                d1["Convex Shape"], d1["Irregularity"],
                d2["Convex Shape"], d2["Irregularity"]
            )
            d_m, t2_stat, f_stat, p_val_t2 = hotelling_t2_2samp(
                d1["Convex Shape"], d1["Irregularity"],
                d2["Convex Shape"], d2["Irregularity"]
            )
            p_ks_str = f"{p_val_ks:.2e}" if (not np.isnan(p_val_ks) and p_val_ks < 0.001) else (round(p_val_ks, 4) if not np.isnan(p_val_ks) else "N/A")
            p_t2_str = f"{p_val_t2:.2e}" if (not np.isnan(p_val_t2) and p_val_t2 < 0.001) else (round(p_val_t2, 4) if not np.isnan(p_val_t2) else "N/A")
            ks_rows.append({
                "Muestra 1": f1.replace(".h5", "").replace(".hdf5", ""),
                "Muestra 2": f2.replace(".h5", "").replace(".hdf5", ""),
                "Estadístico D (K-S 2D)": round(d_stat, 4) if not np.isnan(d_stat) else "N/A",
                "p-valor (K-S 2D)": p_ks_str,
                "Sig. K-S 2D (α=0.05)": "Sí" if (not np.isnan(p_val_ks) and p_val_ks < 0.05) else "No",
                "Dist. Mahalanobis (D_M)": round(d_m, 4) if not np.isnan(d_m) else "N/A",
                "Estadístico T² (Hotelling)": round(t2_stat, 2) if not np.isnan(t2_stat) else "N/A",
                "p-valor (Hotelling T²)": p_t2_str,
                "Sig. Hotelling T² (α=0.05)": "Sí" if (not np.isnan(p_val_t2) and p_val_t2 < 0.05) else "No"
            })
    df_ks = pd.DataFrame(ks_rows)

    # D. Generación de Gráficos Temporales
    img_scatter_path = os.path.join(root_directory, "temp_scatter.png")
    img_kde_path = os.path.join(root_directory, "temp_kde.png")

    plt.figure(figsize=(10, 6))
    for filename, df in file_dfs.items():
        plt.scatter(
            df["Convex Shape"], df["Irregularity"], 
            label=filename.replace(".h5", "").replace(".hdf5", ""), 
            alpha=0.6, edgecolors='none', s=25
        )
    plt.xlabel("Convex Shape", fontsize=12)
    plt.ylabel("Irregularity = atan(Shape / Convex Shape)", fontsize=12)
    plt.title("Irregularity vs. Convex Shape por Archivo", fontsize=14, fontweight='bold')
    plt.grid(True, linestyle="--", alpha=0.5)
    plt.legend(title="Archivos", bbox_to_anchor=(1.05, 1), loc='upper left')
    plt.tight_layout()
    plt.savefig(img_scatter_path, dpi=200, bbox_inches='tight')
    plt.close()

    num_samples = len(file_dfs)
    cols = 2 if num_samples > 1 else 1
    rows = math.ceil(num_samples / cols)
    fig, axes = plt.subplots(rows, cols, figsize=(6 * cols, 5 * rows), sharex=True, sharey=True)
    if num_samples == 1:
        axes = np.array([axes])
    axes_flat = axes.flatten()

    for idx, (filename, df) in enumerate(file_dfs.items()):
        ax = axes_flat[idx]
        clean_name = filename.replace(".h5", "").replace(".hdf5", "")
        sns.kdeplot(
            data=df, x="Convex Shape", y="Irregularity",
            ax=ax, cmap="Blues", fill=True, thresh=0.05, levels=10
        )
        ax.scatter(df["Convex Shape"], df["Irregularity"], s=6, alpha=0.15, color='darkblue')
        ax.set_title(f"Muestra: {clean_name} (n={len(df)})", fontsize=11, fontweight='bold')
        ax.set_xlabel("Convex Shape", fontsize=10)
        ax.set_ylabel("Irregularity = atan(Shape / Convex Shape)", fontsize=10)
        ax.grid(True, linestyle="--", alpha=0.4)
        ax.set_xlim(0, 1)
        ax.set_ylim(0, 0.85)

    for idx in range(num_samples, len(axes_flat)):
        axes_flat[idx].axis('off')

    plt.suptitle("Mapas de Densidad Kernel 2D (Irregularity vs. Convex Shape)", fontsize=14, fontweight='bold', y=0.99)
    plt.tight_layout()
    plt.savefig(img_kde_path, dpi=200, bbox_inches='tight')
    plt.close()

    # E. Guardar en Excel con openpyxl
    excel_path = os.path.join(root_directory, output_excel)
    with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
        final_excel_df.to_excel(writer, sheet_name="Datos Poros")
        df_summary.to_excel(writer, sheet_name="Estadísticas Descriptivas", index=False)
        df_ks.to_excel(writer, sheet_name="Prueba K-S 2D", index=False)

    wb = openpyxl.load_workbook(excel_path)
    ws_chart = wb.create_sheet(title="Gráficos")
    ws_chart.cell(row=2, column=2, value="1. Gráfico de Dispersión Global (Irregularity vs. Convex Shape)").font = openpyxl.styles.Font(bold=True, size=12)
    img_scatter = Image(img_scatter_path)
    ws_chart.add_image(img_scatter, "B4")

    ws_chart.cell(row=38, column=2, value="2. Mapas de Densidad Kernel 2D (KDE 2D) - Mosaico Comparativo").font = openpyxl.styles.Font(bold=True, size=12)
    img_kde = Image(img_kde_path)
    ws_chart.add_image(img_kde, "B40")
    wb.save(excel_path)

    # F. Generar e Insertar Análisis Multivariado p-Dimensional
    if len(file_dfs) >= 2:
        multivariate_df = pd.concat(
            [
                df.assign(Muestra=os.path.splitext(filename)[0])
                for filename, df in file_dfs.items()
            ],
            ignore_index=True,
        )
        try:
            summary_mv, ranking_pivot_mv, detail_mv = run_multivariate_analysis(multivariate_df)
            add_multivariate_tab_to_excel(excel_path, summary_mv, ranking_pivot_mv, detail_mv)
        except ValueError as error:
            print(f"No se pudo generar el análisis multivariado: {error}")

    # Limpiar imágenes temporales
    if os.path.exists(img_scatter_path):
        os.remove(img_scatter_path)
    if os.path.exists(img_kde_path):
        os.remove(img_kde_path)

    print(f"\n¡Proceso completado exitosamente!")
    print(f"📊 Archivo Excel generado: {excel_path}")


# ==========================================
# PUNTO DE EJECUCIÓN
# ==========================================
if __name__ == "__main__":
    CARPETA_RAIZ = r"H:\Otros ordenadores\Microscopio\Silvia EEUU\Porosidad de los mosaicos 4"
    process_all_mosaics(CARPETA_RAIZ)