import os
import threading
import tkinter as tk
import tkinter.filedialog as fd
import tkinter.messagebox as messagebox
from tkinter import ttk

import h5py
import openpyxl as opxl
from PIL import Image

import matplotlib

matplotlib.use("TkAgg")
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg  # noqa: E402
from matplotlib.figure import Figure  # noqa: E402
from matplotlib.patches import Patch  # noqa: E402

# Shapes/sizes used to segment pores (see proc_mosaic.py). Per user decision, the
# ellipse minor diameter ("emd") sheets are used as the single size metric for all shapes.
SHAPES = ["circ", "MLcirc", "shpless", "elongated"]
SIZES = ["S", "M", "L", "XL"]
SIZE_SHEET_PREFIX = "emd"
ROW_HEIGHT = 28


def spectrum_tab(self):
    # State
    self.spectrum_stats_path = None
    self.spectrum_mosaics = []  # list of dicts, order defines plot/list order
    self.spectrum_aspect = 1.0 / 6.0
    self.spectrum_loading_var = tk.StringVar(value="")
    self._spectrum_drag_index = None
    self._spectrum_redraw_after_id = None

    # Left controls
    self.spectrum_frame_controls = ttk.Frame(self.spectrum_frame)
    self.spectrum_frame_controls.pack(side=tk.LEFT, fill=tk.Y, padx=5, pady=5)

    self.spectrum_load_btn = ttk.Button(
        self.spectrum_frame_controls,
        text="Load Global_Pore_Stats.xlsx",
        command=lambda: _spectrum_load_stats(self),
    )
    self.spectrum_load_btn.pack(pady=(0, 8), fill=tk.X)

    # Aspect ratio control (V/H of each individual chart)
    aspect_frame = ttk.Frame(self.spectrum_frame_controls)
    aspect_frame.pack(pady=(0, 8), fill=tk.X)
    self.spectrum_aspect_var = tk.StringVar(value=_format_aspect(self.spectrum_aspect))
    ttk.Button(
        aspect_frame, text="-", width=2, command=lambda: _spectrum_adjust_aspect(self, 1 / 1.25)
    ).pack(side=tk.LEFT)
    ttk.Label(
        aspect_frame, textvariable=self.spectrum_aspect_var, anchor="center", width=10
    ).pack(side=tk.LEFT, expand=True, fill=tk.X)
    ttk.Button(
        aspect_frame, text="+", width=2, command=lambda: _spectrum_adjust_aspect(self, 1.25)
    ).pack(side=tk.LEFT)

    self.spectrum_loading_label = ttk.Label(
        self.spectrum_frame_controls, textvariable=self.spectrum_loading_var, foreground="gray"
    )
    self.spectrum_loading_label.pack(pady=(0, 4), fill=tk.X)

    ttk.Label(self.spectrum_frame_controls, text="Mosaics:").pack(anchor="w")

    # Scrollable, re-orderable list of mosaics
    list_container = ttk.Frame(self.spectrum_frame_controls)
    list_container.pack(fill=tk.BOTH, expand=True)

    self.spectrum_list_canvas = tk.Canvas(list_container, width=220, highlightthickness=0)
    list_scroll = ttk.Scrollbar(
        list_container, orient="vertical", command=self.spectrum_list_canvas.yview
    )
    self.spectrum_list_canvas.configure(yscrollcommand=list_scroll.set)
    self.spectrum_list_canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
    list_scroll.pack(side=tk.RIGHT, fill=tk.Y)

    self.spectrum_list_frame = ttk.Frame(self.spectrum_list_canvas)
    self.spectrum_list_window = self.spectrum_list_canvas.create_window(
        (0, 0), window=self.spectrum_list_frame, anchor="nw"
    )
    self.spectrum_list_frame.bind(
        "<Configure>",
        lambda e: self.spectrum_list_canvas.configure(
            scrollregion=self.spectrum_list_canvas.bbox("all")
        ),
    )
    self.spectrum_list_canvas.bind(
        "<Configure>",
        lambda e: self.spectrum_list_canvas.itemconfigure(self.spectrum_list_window, width=e.width),
    )

    # Right side: stacked column charts
    self.spectrum_frame_display = ttk.Frame(self.spectrum_frame)
    self.spectrum_frame_display.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True, padx=5, pady=5)

    self.spectrum_canvas_holder = tk.Canvas(
        self.spectrum_frame_display, background="white", highlightthickness=0
    )
    display_v_scroll = ttk.Scrollbar(
        self.spectrum_frame_display, orient="vertical", command=self.spectrum_canvas_holder.yview
    )
    self.spectrum_canvas_holder.configure(yscrollcommand=display_v_scroll.set)
    self.spectrum_canvas_holder.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
    display_v_scroll.pack(side=tk.RIGHT, fill=tk.Y)

    self.spectrum_figure = Figure(figsize=(8, 4), dpi=100)
    self.spectrum_canvas_agg = FigureCanvasTkAgg(
        self.spectrum_figure, master=self.spectrum_canvas_holder
    )
    self.spectrum_canvas_widget = self.spectrum_canvas_agg.get_tk_widget()
    self.spectrum_canvas_window = self.spectrum_canvas_holder.create_window(
        (0, 0), window=self.spectrum_canvas_widget, anchor="nw"
    )

    self.spectrum_canvas_holder.bind("<Configure>", lambda e: _spectrum_on_holder_resize(self))

    _spectrum_redraw(self)


# --------------------------------------------------------------------------------------
# Aspect ratio controls
# --------------------------------------------------------------------------------------
def _format_aspect(value):
    if value <= 0:
        return "V/H = 0"
    return f"V/H = 1/{1.0 / value:.1f}"


def _spectrum_adjust_aspect(self, factor):
    new_value = self.spectrum_aspect * factor
    new_value = max(1.0 / 30.0, min(1.0, new_value))
    self.spectrum_aspect = new_value
    self.spectrum_aspect_var.set(_format_aspect(new_value))
    _spectrum_redraw(self)


# --------------------------------------------------------------------------------------
# Loading Global_Pore_Stats.xlsx and per-mosaic shape/size spreadsheets
# --------------------------------------------------------------------------------------
def _spectrum_load_stats(self):
    path = fd.askopenfilename(
        title="Select Global_Pore_Stats.xlsx",
        filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
    )
    if not path:
        return

    self.spectrum_load_btn.config(state=tk.DISABLED)
    self.spectrum_loading_var.set("Loading...")

    def worker():
        try:
            entries = _spectrum_read_global_stats(path)
            base_dir = os.path.dirname(path)
            for entry in entries:
                entry.update(_spectrum_compute_mosaic_metrics(base_dir, entry["name"]))
        except Exception as e:
            self.root.after(0, lambda: _spectrum_load_failed(self, str(e)))
            return
        self.root.after(0, lambda: _spectrum_load_done(self, path, entries))

    threading.Thread(target=worker, daemon=True).start()


def _spectrum_load_failed(self, message):
    self.spectrum_load_btn.config(state=tk.NORMAL)
    self.spectrum_loading_var.set("")
    messagebox.showerror("Error", f"Failed to load Global_Pore_Stats.xlsx: {message}")


def _spectrum_load_done(self, path, entries):
    self.spectrum_stats_path = path
    self.spectrum_mosaics = [
        {**entry, "enabled": tk.BooleanVar(value=True)} for entry in entries
    ]
    self.spectrum_load_btn.config(state=tk.NORMAL)
    self.spectrum_loading_var.set(f"Loaded {len(self.spectrum_mosaics)} mosaic(s).")
    _spectrum_rebuild_list_ui(self)
    _spectrum_redraw(self)


def _spectrum_read_global_stats(path):
    """Read mosaic names and global stats from Global_Pore_Stats.xlsx."""
    wb = opxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    if not rows:
        return []

    header = [str(h) if h is not None else "" for h in rows[0]]
    col = {name: i for i, name in enumerate(header)}

    required = [
        "Mosaic Name",
        "Porosity",
        "Percentage of pores <= 50μm",
        "Percentage of pores > 50μm",
    ]
    missing = [name for name in required if name not in col]
    if missing:
        raise ValueError(f"Missing expected column(s): {', '.join(missing)}")

    entries = []
    for row in rows[1:]:
        if not row or row[col["Mosaic Name"]] is None:
            continue
        entries.append(
            {
                "name": str(row[col["Mosaic Name"]]),
                "porosity": float(row[col["Porosity"]] or 0.0),
                "pct_le50": float(row[col["Percentage of pores <= 50μm"]] or 0.0),
                "pct_gt50": float(row[col["Percentage of pores > 50μm"]] or 0.0),
            }
        )
    return entries


def _spectrum_compute_mosaic_metrics(base_dir, name):
    """Compute area-fraction / average-area / count metrics for a single mosaic."""
    mosaic_dir = os.path.join(base_dir, name)
    tiff_path = os.path.join(mosaic_dir, name + ".tiff")
    h5_path = os.path.join(mosaic_dir, name + ".h5")

    image_area_um2 = None
    if os.path.exists(tiff_path):
        try:
            with Image.open(tiff_path) as img:
                width_px, height_px = img.size
            calibration = None
            if os.path.exists(h5_path):
                with h5py.File(h5_path, "r") as f:
                    calibration = f.attrs.get("pixel_calibration_px_per_um")
            if calibration:
                image_area_um2 = (width_px * height_px) / (float(calibration) ** 2)
        except Exception:
            image_area_um2 = None

    area_fraction = {}
    avg_area = {}
    count = {}

    for shape in SHAPES:
        shape_path = os.path.join(mosaic_dir, f"{name}_{shape}.xlsx")
        if not os.path.exists(shape_path):
            for size in SIZES:
                area_fraction[(shape, size)] = 0.0
                avg_area[(shape, size)] = 0.0
                count[(shape, size)] = 0.0
            continue

        wb = opxl.load_workbook(shape_path, data_only=True, read_only=True)
        try:
            for size in SIZES:
                sheet_name = f"{SIZE_SHEET_PREFIX}{size}"
                key = (shape, size)
                if sheet_name not in wb.sheetnames:
                    area_fraction[key] = 0.0
                    avg_area[key] = 0.0
                    count[key] = 0.0
                    continue

                ws = wb[sheet_name]
                total_area = 0.0
                non_edge_areas = []
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or row[0] is None:
                        continue
                    area = float(row[2] or 0.0)
                    total_area += area
                    is_edge = str(row[1]).strip().lower() == "true"
                    if not is_edge:
                        non_edge_areas.append(area)

                area_fraction[key] = total_area / image_area_um2 if image_area_um2 else 0.0
                avg_area[key] = (
                    sum(non_edge_areas) / len(non_edge_areas) if non_edge_areas else 0.0
                )
                count[key] = float(len(non_edge_areas))
        finally:
            wb.close()

    return {"area_fraction": area_fraction, "avg_area": avg_area, "count": count}


# --------------------------------------------------------------------------------------
# Re-orderable / checkable mosaic list
# --------------------------------------------------------------------------------------
def _spectrum_rebuild_list_ui(self):
    for child in self.spectrum_list_frame.winfo_children():
        child.destroy()

    for index, entry in enumerate(self.spectrum_mosaics):
        row = tk.Frame(self.spectrum_list_frame, height=ROW_HEIGHT)
        row.pack(fill=tk.X)
        row.pack_propagate(False)

        handle = ttk.Label(row, text="⋮⋮", cursor="fleur")
        handle.pack(side=tk.LEFT, padx=(2, 6))
        handle.bind("<ButtonPress-1>", lambda e, i=index: _spectrum_start_drag(self, i))
        handle.bind("<B1-Motion>", lambda e: _spectrum_drag_motion(self, e))
        handle.bind("<ButtonRelease-1>", lambda e: _spectrum_end_drag(self, e))

        check = ttk.Checkbutton(
            row, variable=entry["enabled"], command=lambda: _spectrum_redraw(self)
        )
        check.pack(side=tk.LEFT, padx=(0, 4))

        ttk.Label(row, text=entry["name"]).pack(side=tk.LEFT, fill=tk.X, expand=True)

    self.spectrum_list_canvas.update_idletasks()
    self.spectrum_list_canvas.configure(scrollregion=self.spectrum_list_canvas.bbox("all"))


def _spectrum_start_drag(self, index):
    self._spectrum_drag_index = index


def _spectrum_drag_motion(self, event):
    if self._spectrum_drag_index is None:
        return
    rel_y = event.y_root - self.spectrum_list_frame.winfo_rooty()
    target = max(0, min(len(self.spectrum_mosaics) - 1, int(rel_y // ROW_HEIGHT)))
    if target != self._spectrum_drag_index:
        entry = self.spectrum_mosaics.pop(self._spectrum_drag_index)
        self.spectrum_mosaics.insert(target, entry)
        self._spectrum_drag_index = target
        _spectrum_rebuild_list_ui(self)


def _spectrum_end_drag(self, event):
    if self._spectrum_drag_index is not None:
        self._spectrum_drag_index = None
        _spectrum_redraw(self)


# --------------------------------------------------------------------------------------
# Data Preparation & Chart Rendering (1 Chart per Mosaic)
# --------------------------------------------------------------------------------------
def _spectrum_prepare_data(mosaics):
    """
    Builds the array of 51 column labels and normalized values (0 to 1) for each mosaic.
    Categories:
      1. General stats (3)
      2. Area fraction (16)
      3. Average pore area (16, normalized to max across loaded mosaics)
      4. Pore count (16, normalized to max across loaded mosaics)
    """
    combos = [(shape, size) for shape in SHAPES for size in SIZES]

    # Max values for normalization of avg area & count across all loaded mosaics
    max_avg_area = {c: 0.0 for c in combos}
    max_count = {c: 0.0 for c in combos}

    for m in mosaics:
        for c in combos:
            if m["avg_area"].get(c, 0.0) > max_avg_area[c]:
                max_avg_area[c] = m["avg_area"][c]
            if m["count"].get(c, 0.0) > max_count[c]:
                max_count[c] = m["count"][c]

    labels = []
    # 1. General stats
    labels.extend(["Porosity", "% ≤ 50μm", "% > 50μm"])

    # 2. Area fraction
    for shape, size in combos:
        labels.append(f"AF {shape}-{size}")

    # 3. Avg area
    for shape, size in combos:
        labels.append(f"AvgA {shape}-{size}")

    # 4. Count
    for shape, size in combos:
        labels.append(f"Cnt {shape}-{size}")

    data_by_mosaic = {}
    for m in mosaics:
        vals = []

        # General stats (ensure values are scale 0..1 if given as percentages > 1)
        p = m["porosity"]
        p = p / 100.0 if p > 1.0 else p
        le = m["pct_le50"]
        le = le / 100.0 if le > 1.0 else le
        gt = m["pct_gt50"]
        gt = gt / 100.0 if gt > 1.0 else gt
        vals.extend([max(0.0, min(1.0, p)), max(0.0, min(1.0, le)), max(0.0, min(1.0, gt))])

        # Area fraction (16 values)
        for c in combos:
            af = m["area_fraction"].get(c, 0.0)
            af = af / 100.0 if af > 1.0 else af
            vals.append(max(0.0, min(1.0, af)))

        # Avg area (16 values normalized between 0 and 1)
        for c in combos:
            val = m["avg_area"].get(c, 0.0)
            norm = (val / max_avg_area[c]) if max_avg_area[c] > 0 else 0.0
            vals.append(max(0.0, min(1.0, norm)))

        # Pore count (16 values normalized between 0 and 1)
        for c in combos:
            val = m["count"].get(c, 0.0)
            norm = (val / max_count[c]) if max_count[c] > 0 else 0.0
            vals.append(max(0.0, min(1.0, norm)))

        data_by_mosaic[m["name"]] = vals

    return labels, data_by_mosaic


def _spectrum_on_holder_resize(self):
    if self._spectrum_redraw_after_id is not None:
        self.root.after_cancel(self._spectrum_redraw_after_id)
    self._spectrum_redraw_after_id = self.root.after(150, lambda: _spectrum_redraw(self))


def _spectrum_apply_figure_size(self, rows):
    self.spectrum_canvas_holder.update_idletasks()
    width_px = self.spectrum_canvas_holder.winfo_width()
    if width_px <= 1:
        width_px = 800
    dpi = self.spectrum_figure.get_dpi()
    width_in = width_px / dpi
    height_in = max(width_in * self.spectrum_aspect * rows, 1.2)
    self.spectrum_figure.set_size_inches(width_in, height_in, forward=True)
    self.spectrum_canvas_widget.configure(width=width_px, height=int(height_in * dpi))


def _spectrum_update_scrollregion(self):
    self.spectrum_canvas_holder.update_idletasks()
    self.spectrum_canvas_holder.itemconfigure(
        self.spectrum_canvas_window, width=self.spectrum_canvas_holder.winfo_width()
    )
    self.spectrum_canvas_holder.configure(
        scrollregion=self.spectrum_canvas_holder.bbox(self.spectrum_canvas_window)
    )


def _spectrum_show_message(self, message):
    fig = self.spectrum_figure
    ax = fig.add_subplot(111)
    ax.text(0.5, 0.5, message, ha="center", va="center", fontsize=10)
    ax.axis("off")
    _spectrum_apply_figure_size(self, 1)
    self.spectrum_canvas_agg.draw()
    _spectrum_update_scrollregion(self)


def _spectrum_redraw(self):
    if self._spectrum_redraw_after_id is not None:
        self.root.after_cancel(self._spectrum_redraw_after_id)
        self._spectrum_redraw_after_id = None

    self.spectrum_figure.clear()

    if not self.spectrum_mosaics:
        _spectrum_show_message(self, "Load Global_Pore_Stats.xlsx to begin.")
        return

    visible_mosaics = [m for m in self.spectrum_mosaics if m["enabled"].get()]
    if not visible_mosaics:
        _spectrum_show_message(self, "No mosaics selected for display.")
        return

    labels, data_by_mosaic = _spectrum_prepare_data(self.spectrum_mosaics)
    n_mosaics = len(visible_mosaics)

    # Create 1 subplot per visible mosaic (stacked vertically)
    axes = self.spectrum_figure.subplots(n_mosaics, 1, sharex=True)
    if n_mosaics == 1:
        axes = [axes]

    # Distinguish categories by color: Blue (General), Green (Area Fraction), Orange (Avg Area), Purple (Count)
    bar_colors = ["#1f77b4"] * 3 + ["#2ca02c"] * 16 + ["#ff7f0e"] * 16 + ["#9467bd"] * 16
    x = list(range(len(labels)))

    for idx, m in enumerate(visible_mosaics):
        ax = axes[idx]
        vals = data_by_mosaic.get(m["name"], [0.0] * len(labels))
        ax.bar(x, vals, color=bar_colors, width=0.8)
        ax.set_ylim(0, 1.05)
        ax.set_title(f"Mosaic: {m['name']}", fontsize=9, loc="left", fontweight="bold")
        ax.tick_params(axis="y", labelsize=7)
        ax.margins(x=0.01)

        # Dashed vertical separators between category groups
        ax.axvline(2.5, color="gray", linestyle="--", linewidth=0.8, alpha=0.6)
        ax.axvline(18.5, color="gray", linestyle="--", linewidth=0.8, alpha=0.6)
        ax.axvline(34.5, color="gray", linestyle="--", linewidth=0.8, alpha=0.6)

        # Tick labels on bottom axis
        if idx == n_mosaics - 1:
            ax.set_xticks(x)
            ax.set_xticklabels(labels, rotation=90, ha="center", fontsize=6.5)
            ax.tick_params(axis="x", labelbottom=True)
        else:
            ax.set_xticks(x)
            ax.tick_params(axis="x", labelbottom=False)

    # Figure legend for metric groups
    legend_elements = [
        Patch(facecolor="#1f77b4", label="General Stats"),
        Patch(facecolor="#2ca02c", label="Area Fraction"),
        Patch(facecolor="#ff7f0e", label="Avg Area (norm)"),
        Patch(facecolor="#9467bd", label="Pore Count (norm)"),
    ]
    self.spectrum_figure.legend(
        handles=legend_elements,
        loc="upper center",
        ncol=4,
        fontsize=8,
        bbox_to_anchor=(0.5, 1.01),
    )
    self.spectrum_figure.tight_layout(rect=[0, 0, 1, 0.95])

    _spectrum_apply_figure_size(self, n_mosaics)
    self.spectrum_canvas_agg.draw()
    _spectrum_update_scrollregion(self)